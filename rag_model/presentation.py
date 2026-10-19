"""Checked pitch items -> slide deck (presentation/template/raglette.typ, Kanton Bern design).

    deck_data(items, ...)   our items in the deck's JSON format
    slide_map(data)         what each slide shows (so the UI can delete the slide of an item)
    render(data, "pdf")     the PDF as bytes
    render(data, "png")     one PNG per slide, for the preview in the UI
    pptx(data)              PowerPoint: every slide as a picture of the Typst slide (looks exactly like the PDF),
                            with the slide's text in the speaker notes to copy from

Slides: title · at a glance · one slide per item (grouped by category) · contacts. Without a named contact
the source page is the contact (its website and a link). The fixed labels follow Settings → Result → language.
"""
import io
import json
import textwrap

import typst
from pptx import Presentation
from pptx.util import Cm

from rag_model import config
from rag_model.search import ROOT

TEMPLATE = ROOT / "presentation" / "template"
MAX_CONTACTS = 3          # per slide, more do not fit next to the text
MAX_QUOTE = 320           # characters of the quote shown on the slide (the full quote stays in the evidence)
CATEGORIES = ["research", "infrastructure", "programme", "network", "education"]   # also the slide order

LABELS = {
    "English": {
        "kicker": "Business Case", "subtitle": "What Bern brings to the table", "speaker": "Speaker",
        "contacts": "Contacts", "source": "Source", "overview": "At a glance",
        "overview_cols": ["Category", "Partner", "Organisation"],
        "what": "What it is", "why": "Why it matters for {company}", "facts": "Key facts",
        "contacts_title": "Your contacts in Bern", "contacts_cols": ["Name", "Role", "Topic", "Contact"],
        "categories": {"research": "Research partner", "infrastructure": "Infrastructure",
                       "programme": "Programme & funding", "network": "Network", "education": "Education & talent"},
    },
    "German": {
        "kicker": "Business Case", "subtitle": "Was Bern zu bieten hat", "speaker": "Referent/-in",
        "contacts": "Ansprechpersonen", "source": "Quelle", "overview": "Auf einen Blick",
        "overview_cols": ["Kategorie", "Partner", "Organisation"],
        "what": "Worum es geht", "why": "Warum für {company}", "facts": "Fakten",
        "contacts_title": "Ihre Ansprechpersonen in Bern", "contacts_cols": ["Name", "Funktion", "Thema", "Kontakt"],
        "categories": {"research": "Forschungspartner", "infrastructure": "Infrastruktur",
                       "programme": "Programme & Förderung", "network": "Netzwerk", "education": "Bildung & Talente"},
    },
    "French": {
        "kicker": "Business Case", "subtitle": "Ce que Berne a à offrir", "speaker": "Intervenant·e",
        "contacts": "Personnes de contact", "source": "Source", "overview": "En un coup d'œil",
        "overview_cols": ["Catégorie", "Partenaire", "Organisation"],
        "what": "De quoi s'agit-il", "why": "Pourquoi pour {company}", "facts": "Faits",
        "contacts_title": "Vos contacts à Berne", "contacts_cols": ["Nom", "Fonction", "Thème", "Contact"],
        "categories": {"research": "Partenaire de recherche", "infrastructure": "Infrastructure",
                       "programme": "Programmes & soutien", "network": "Réseau", "education": "Formation & talents"},
    },
}


def category(value: str) -> str:
    """One of CATEGORIES; free text from older items is mapped by its words."""
    v = (value or "").strip().lower()
    if v in CATEGORIES:
        return v
    for key, words in (("infrastructure", ("infra", "platform", "facilit", "lab equipment")),
                       ("programme", ("program", "fund", "grant", "support", "accelerat", "incubat", "competition")),
                       ("network", ("network", "cluster", "association", "consortium", "hub")),
                       ("education", ("educat", "talent", "course", "study", "training", "degree"))):
        if any(w in v for w in words):
            return key
    return "research"


def _contacts(item: dict) -> list[dict]:
    out = []
    for c in item.get("contacts") or []:
        if not c.get("name"):
            continue
        entry = {"name": c["name"]}
        for ours, theirs in (("role", "function"), ("email", "email"), ("phone", "phone")):
            if c.get(ours):
                entry[theirs] = c[ours]
        out.append(entry)
    if not out:
        out = [{"name": s["website"], "url": s["url"], "url_label": textwrap.shorten(s.get("title") or s["url"], 60)}
               for s in item.get("sources", [])[:1]]
    return out[:MAX_CONTACTS]


DECK = {"company": "", "subtitle": "", "kicker": "", "author": "", "organizer": "",
        "overview": True, "contacts_slide": True}   # the presentation settings of a conversation


def deck_data(items: list[dict], company: str, author: str, organizer: str,
              subtitle: str = "", kicker: str = "", overview: bool = True, contacts_slide: bool = True) -> dict:
    """subtitle, kicker: the title slide's subtitle and the small red label above the title; empty = the default
    in the deck language. overview, contacts_slide: show the "at a glance" and the contacts slide."""
    labels = dict(LABELS.get(config.load()["brief"]["language"], LABELS["English"]))
    labels.update({k: v.strip() for k, v in (("subtitle", subtitle), ("kicker", kicker)) if v.strip()})
    company = company.strip() or "Company"
    rows = sorted(items, key=lambda it: CATEGORIES.index(category(it.get("category", ""))))
    return {
        "company": company, "author": author.strip(), "organizer": organizer.strip(),
        "categories": CATEGORIES, "show_overview": overview, "show_contacts": contacts_slide,
        "labels": {**{k: v for k, v in labels.items() if k != "categories"}, "why": labels["why"].format(company=company)},
        "items": [{
            "title": it.get("title", ""),
            "category": category(it.get("category", "")),
            "category_label": labels["categories"][category(it.get("category", ""))],
            "organisation": it.get("organisation") or "",
            "summary": " ".join((it.get("summary") or "").split()),
            "relevance": " ".join((it.get("relevance") or "").split()),
            "facts": [f for f in it.get("facts") or [] if isinstance(f, str) and f.strip()][:3],
            "evidence": textwrap.shorten(it.get("evidence") or "", MAX_QUOTE, placeholder=" …"),
            "contacts": _contacts(it),
            "sources": [{"label": s.get("website") or s["url"], "url": s["url"]} for s in it.get("sources", [])],
        } for it in rows],
    }


def slide_map(data: dict) -> list[dict]:
    """One entry per slide, in the order raglette.typ draws them."""
    items = data["items"]
    people = data["show_contacts"] and any("url" not in c for it in items for c in it["contacts"])
    return ([{"kind": "title"}] + ([{"kind": "overview"}] if items and data["show_overview"] else [])
            + [{"kind": "item", "title": it["title"]} for it in items] + ([{"kind": "contacts"}] if people else []))


def defaults() -> dict:
    """The default title-slide texts in the deck language, shown as placeholders in the UI."""
    labels = LABELS.get(config.load()["brief"]["language"], LABELS["English"])
    return {"subtitle": labels["subtitle"], "kicker": labels["kicker"]}


def render(data: dict, fmt: str = "pdf", ppi: int = 72) -> bytes | list[bytes]:
    out = typst.compile(TEMPLATE / "raglette.typ", root=TEMPLATE, format=fmt, ppi=ppi if fmt == "png" else None,
                        font_paths=[TEMPLATE / "assets" / "fonts"], ignore_system_fonts=True,
                        sys_inputs={"json": json.dumps(data, ensure_ascii=False)})
    return out if isinstance(out, list) or fmt == "pdf" else [out]


PPTX_PPI = 200   # sharp on a projector, a few MB per deck


def _notes(data: dict, info: dict) -> str:
    """The text of one slide, for the speaker notes."""
    L = data["labels"]
    if info["kind"] == "title":
        return "\n".join(t for t in (data["company"], L["subtitle"], f"{L['speaker']}: {data['author']}", data["organizer"]) if t)
    if info["kind"] == "overview":
        return "\n".join(f"{it['category_label']}: {it['title']} ({it['organisation']})" for it in data["items"])
    if info["kind"] == "contacts":
        return "\n".join(f"{c['name']}, {c.get('function', '')}, {c.get('email') or c.get('phone') or ''} – {it['title']}"
                         for it in data["items"] for c in it["contacts"] if "url" not in c)
    it = next(i for i in data["items"] if i["title"] == info["title"])
    lines = [it["title"], " · ".join(t for t in (it["category_label"], it["organisation"]) if t), "",
             f"{L['what']}: {it['summary']}"]
    if it["relevance"]:
        lines.append(f"{L['why']}: {it['relevance']}")
    lines += [f"– {f}" for f in it["facts"]]
    if it["evidence"]:
        lines += ["", f"«{it['evidence']}»"]
    lines += ["", f"{L['source']}: " + " · ".join(s["url"] for s in it["sources"])]
    return "\n".join(lines)


def pptx(data: dict) -> bytes:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Cm(33.867), Cm(19.05)
    blank = deck.slide_layouts[6]
    for png, info in zip(render(data, "png", ppi=PPTX_PPI), slide_map(data)):
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(png), 0, 0, width=deck.slide_width, height=deck.slide_height)
        slide.notes_slide.notes_text_frame.text = _notes(data, info)
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


if __name__ == "__main__":   # writes raglette-example.json from the latest agent trace, to compile the deck by hand
    import glob
    import os
    trace = json.load(open(max(glob.glob(str(ROOT / "data" / "traces" / "*.json")), key=os.path.getmtime)))
    data = deck_data(trace.get("items") or [], "Example AG", "Anna Muster", "Standortförderung Kanton Bern")
    (TEMPLATE / "raglette-example.json").write_text(json.dumps(data, ensure_ascii=False, indent=2))
    print(f"{len(data['items'])} items -> {TEMPLATE / 'raglette-example.json'}")
