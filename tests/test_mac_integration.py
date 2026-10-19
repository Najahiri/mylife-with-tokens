"""Exercise the mac UI's backend contracts with local data and mocked model calls.

Run with the project environment: uv run python -m unittest discover -s tests -p test_mac_integration.py
"""

from contextlib import ExitStack, closing
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient
from haystack.dataclasses import ChatMessage, ToolCall
from pptx import Presentation

from api import main as api
from rag_model import agent, config, llm, pipeline_log, policy, settings


class MacIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.root = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        (self.root / "data").mkdir()
        with closing(sqlite3.connect(self.root / "data/index.sqlite")) as db:
            db.execute("CREATE TABLE pages (source TEXT)")
            db.execute("CREATE TABLE chunks (id TEXT)")
            db.commit()
        for owner, name, value in (
            (api, "ROOT", self.root),
            (api, "Searcher", MagicMock()),
            (api.indexer, "load", MagicMock()),
            (api, "state", {}),
            (api, "sessions", {}),
            (api, "busy", {}),
            (agent, "SESSIONS", self.root / "sessions"),
            (llm, "SETTINGS_FILE", self.root / "settings.json"),
            (config, "OVERRIDES_FILE", self.root / "overrides.json"),
            (pipeline_log, "LOG_DIR", self.root / "logs"),
            (agent, "write_trace", MagicMock(return_value="test-trace")),
        ):
            self.stack.enter_context(patch.object(owner, name, value))
        self.stack.enter_context(patch.dict(os.environ, {}, clear=True))
        self.client = self.stack.enter_context(TestClient(api.app))

    def test_settings_keep_independent_routes_and_hide_keys(self):
        response = self.client.put(
            "/api/settings",
            json={
                "base_url": llm.KEERA_BASE_URL,
                "api_key": "chat-secret",
                "model_agent": "apertus-v1-5-70b",
                "decision_url": "https://decision.example/",
                "decision_api_key": "decision-secret",
                "decision_model": "our/finetuned-apertus",
            },
        )
        self.assertEqual(response.status_code, 200)
        public = response.json()["llm"]
        self.assertEqual(public["decision_url"], "https://decision.example/decide")
        self.assertEqual(public["model_agent"], "apertus-v1-5-70b")
        self.assertNotIn("api_key", public)
        self.assertNotIn("decision_api_key", public)
        self.assertTrue(self.client.get("/api/health").json()["agent_configured"])
        llm.SETTINGS_FILE.unlink()
        self.assertFalse(self.client.get("/api/health").json()["agent_configured"])

    def test_connection_checks_both_generation_and_decision(self):
        llm.save_settings({"decision_url": "https://decision.example/decide"})
        chat = MagicMock()
        chat.run.return_value = {"replies": [ChatMessage.from_assistant("ok")]}
        with patch.object(settings, "generator", return_value=chat), patch(
            "rag_model.decision.decide", return_value={"prediction": "Option A"}
        ) as decide:
            response = self.client.post("/api/settings/test")
        self.assertEqual(
            [r["role"] for r in response.json()["results"]], ["agent", "decision"]
        )
        decide.assert_called_once()

    def test_history_delete_and_legacy_session_loading(self):
        sid = self.client.post("/api/agent").json()["id"]
        session = api.sessions[sid]
        session.events.append({"type": "user", "text": "Find quantum research"})
        session.trash.append({"title": "Removed partner"})
        session.save()
        loaded = agent.Session.load(agent.SESSIONS / f"{sid}.json")
        self.assertEqual(loaded.trash, session.trash)
        self.assertFalse(loaded.stop.is_set())
        self.assertEqual(
            self.client.get("/api/agent").json()["sessions"][0]["title"],
            "Find quantum research",
        )
        legacy = json.loads((agent.SESSIONS / f"{sid}.json").read_text())
        legacy.pop("updated")
        legacy.pop("trash")
        (agent.SESSIONS / f"{sid}.json").write_text(json.dumps(legacy))
        self.assertEqual(agent.Session.load(agent.SESSIONS / f"{sid}.json").trash, [])
        api.busy[sid].acquire()
        self.assertEqual(self.client.delete(f"/api/agent/{sid}").status_code, 409)
        self.assertTrue(self.client.post(f"/api/agent/{sid}/stop").json()["stopped"])
        self.assertTrue(session.stop.is_set())
        api.busy[sid].release()
        self.assertEqual(self.client.delete(f"/api/agent/{sid}").status_code, 200)
        self.assertFalse((agent.SESSIONS / f"{sid}.json").exists())

    def test_stop_prevents_selected_tool_and_allows_continuation(self):
        session = agent.Session()
        call = ToolCall(
            tool_name="search",
            arguments={"passages": ["Quantum research"]},
            id="call-1",
        )
        chat = MagicMock()
        chat.run.return_value = {
            "replies": [ChatMessage.from_assistant(tool_calls=[call])]
        }
        choice = {
            "action": "search",
            "actions": ["search", "answer"],
            "probabilities": [1.0, 0.0],
        }
        with patch.object(agent, "select_action", return_value=choice), patch.object(
            agent, "generator", return_value=chat
        ), patch.object(agent, "_run_tool") as execute:
            events = agent.turn(session, "Find research", MagicMock())
            for event in events:
                if event["type"] == "tool_call":
                    session.stop.set()
                    break
            remaining = list(events)
        execute.assert_not_called()
        self.assertEqual(remaining, [{"type": "stopped"}])
        self.assertNotIn(agent._call_key(call), agent._completed_calls(session))
        self.assertTrue(session.messages[-1].tool_call_result.error)
        chat.run.return_value = {
            "replies": [ChatMessage.from_assistant("Stopped. No slides were added.")]
        }
        with patch.object(
            agent,
            "select_action",
            return_value={
                "action": "answer",
                "actions": ["answer"],
                "probabilities": [1.0],
            },
        ), patch.object(agent, "generator", return_value=chat):
            continued = list(agent.turn(session, "What happened?", MagicMock()))
        self.assertFalse(session.stop.is_set())
        self.assertTrue(any(e.get("final") for e in continued))

    def test_stop_callback_interrupts_generation(self):
        session = agent.Session()

        def make_chat(*args, on_chunk, **kwargs):
            chat = MagicMock()

            def run(**kwargs):
                session.stop.set()
                on_chunk(None)

            chat.run.side_effect = run
            return chat

        with patch.object(agent, "generator", side_effect=make_chat), patch.object(
            agent,
            "select_action",
            return_value={
                "action": "answer",
                "actions": ["answer"],
                "probabilities": [1.0],
            },
        ):
            events = list(agent.turn(session, "Hello", MagicMock()))
        self.assertEqual(events[-1], {"type": "stopped"})

    def test_restore_is_routed_through_decision_and_text_calls_are_recovered(self):
        session = agent.Session(items=[{"title": "Research partner"}])
        agent.remove_by_user(session, ["Research partner"])
        self.assertEqual(session.items, [])
        state = agent._decision_state(session, "Restore that slide", None)
        with patch.object(
            policy, "decide", return_value={"prediction_index": 0}
        ) as decide:
            result = policy.select_action(state, agent._tools(15))
        self.assertIn("restore_items", result["actions"])
        self.assertLessEqual(len(decide.call_args.args[2]), 8)
        reply, _ = agent._clean(
            ChatMessage.from_assistant('restore_items(titles: ["Research partner"])')
        )
        self.assertFalse(reply.text)
        self.assertEqual(reply.tool_calls[0].tool_name, "restore_items")
        _, info = agent._run_tool(
            session, MagicMock(), "restore_items", reply.tool_calls[0].arguments, 15
        )
        self.assertEqual(info["restored"], 1)
        self.assertEqual(session.items[0]["title"], "Research partner")
        self.assertEqual(session.trash, [])
        self.assertNotIn(
            "uniqueItems", json.dumps([t.parameters for t in agent._tools(15)])
        )

    def test_pdf_and_powerpoint_exports_contain_checked_slide(self):
        item = {
            "title": "Research partner",
            "category": "research",
            "summary": "Quoted research.",
            "evidence": "Quoted research.",
            "sources": [{"url": "https://example.org", "website": "University"}],
        }
        body = {"items": [item], "company": "Example AG"}
        response = self.client.post(
            "/api/presentation/export", json={**body, "format": "pdf"}
        )
        self.assertEqual(response.status_code, 200, response.text[:300])
        self.assertTrue(response.content.startswith(b"%PDF"))
        response = self.client.post(
            "/api/presentation/export", json={**body, "format": "pptx"}
        )
        self.assertEqual(response.status_code, 200, response.text[:300])
        deck = Presentation(io.BytesIO(response.content))
        self.assertEqual(len(deck.slides), 3)
        self.assertIn(
            "Research partner", deck.slides[2].notes_slide.notes_text_frame.text
        )


if __name__ == "__main__":
    unittest.main()
